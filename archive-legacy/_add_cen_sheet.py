"""
Add Sheet 7: CEN Deliverables to CONTROL_PANEL.xlsx
Tracks what exists, what needs packaging, and what goes into the 4 empty phase folders.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

path = r"C:\Users\murau\OneDrive\Stalinis kompiuteris\POC\CONTROL_PANEL.xlsx"
wb = openpyxl.load_workbook(path)

# ── Brand colors ──
DARK_CYAN = "00838F"
CYAN = "00BCD4"
DARK = "263238"
WHITE = "FFFFFF"
GRAY = "9E9E9E"
RED = "D32F2F"
ORANGE = "FF6F00"

header_font = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
header_fill = PatternFill(start_color=DARK_CYAN, end_color=DARK_CYAN, fill_type="solid")
subheader_font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
subheader_fill = PatternFill(start_color=CYAN, end_color=CYAN, fill_type="solid")
section_font = Font(name="Segoe UI", size=11, bold=True, color=DARK_CYAN)
body_font = Font(name="Segoe UI", size=10, color=DARK)
bold_font = Font(name="Segoe UI", size=10, bold=True, color=DARK)
title_font = Font(name="Segoe UI", size=16, bold=True, color=DARK_CYAN)
note_font = Font(name="Segoe UI", size=9, italic=True, color=GRAY)
urgent_font = Font(name="Segoe UI", size=10, bold=True, color=RED)
warn_font = Font(name="Segoe UI", size=10, bold=True, color=ORANGE)
thin_border = Border(
    left=Side(style="thin", color="BDBDBD"),
    right=Side(style="thin", color="BDBDBD"),
    top=Side(style="thin", color="BDBDBD"),
    bottom=Side(style="thin", color="BDBDBD"),
)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center")

pass_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
warn_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")
fail_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
done_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
overdue_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
pending_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")
future_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def cell(ws, r, c, val=None, font=None, fill=None, align=None):
    cl = ws.cell(row=r, column=c, value=val)
    cl.font = font or body_font
    cl.border = thin_border
    cl.alignment = align or wrap
    if fill:
        cl.fill = fill
    return cl


ws = wb.create_sheet("CEN Deliverables")
ws.sheet_properties.tabColor = "FF6F00"

# ── Title ──
ws.merge_cells("A1:H1")
ws["A1"].value = "CEN PARTNERSHIP — DELIVERABLES TRACKER"
ws["A1"].font = title_font
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

# ── Context line ──
ws.merge_cells("A2:H2")
ws["A2"].value = (
    "External Partner: Coherence Partnership Quannex x CEN  |  "
    "Path: Final Thesis/Thesis Work/External Partners/Coherence Partnership Quannex x CEN"
)
ws["A2"].font = note_font

# ═══════════════════════════════════════════════════════════
# SECTION 1: PHASE OVERVIEW
# ═══════════════════════════════════════════════════════════
r = 4
ws.merge_cells(f"A{r}:H{r}")
cell(ws, r, 1, "PHASE OVERVIEW", font=section_font)

r = 5
headers = ["Phase", "Activity", "Date Window", "Status", "Work Done?", "Packaged?", "Delivery Folder", "Notes"]
for c, h in enumerate(headers, 1):
    cl = ws.cell(row=r, column=c, value=h)
    cl.font = header_font
    cl.fill = header_fill
    cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cl.border = thin_border

phases = [
    ("Phase 1", "Kick-off + interview + synthesis", "Mar 9-20", "DONE", "YES", "NO",
     "Deliverables from Quannex/Phase 1", "Interview Mar 20, synthesis complete. FOLDER EMPTY — needs packaging"),
    ("Phase 2", "Scoring + BSC comparison + evidence", "Apr 1-14", "DONE", "YES", "NO",
     "Deliverables from Quannex/Phase 2", "Both questionnaires returned Apr 7. 14-doc evidence package. BSC 5/6. FOLDER EMPTY"),
    ("Phase 3", "Validation session with D + E", "Apr 15-30", "UPCOMING", "NO", "NO",
     "Deliverables from Quannex/Phase 3", "Present coherence portrait. ~Apr 17. Prepare CEN-facing summary"),
    ("Phase 4", "Final report + recommendations", "May", "NOT STARTED", "NO", "NO",
     "Deliverables from Quannex/Phase 4", "Depends on Ch5 + cross-case analysis completion"),
]

for i, (phase, activity, dates, status, work, packaged, folder, notes) in enumerate(phases):
    row = r + 1 + i
    cell(ws, row, 1, phase, font=bold_font, align=center)
    cell(ws, row, 2, activity)
    cell(ws, row, 3, dates, align=center)

    # Status with color
    status_fill = done_fill if status == "DONE" else (warn_fill if status == "UPCOMING" else future_fill)
    status_font = bold_font if status == "DONE" else (warn_font if status == "UPCOMING" else body_font)
    cell(ws, row, 4, status, font=status_font, fill=status_fill, align=center)

    # Work done
    work_fill = done_fill if work == "YES" else future_fill
    cell(ws, row, 5, work, fill=work_fill, align=center)

    # Packaged — this is the gap
    pkg_fill = fail_fill if (work == "YES" and packaged == "NO") else (done_fill if packaged == "YES" else future_fill)
    pkg_font = urgent_font if (work == "YES" and packaged == "NO") else body_font
    cell(ws, row, 6, packaged, font=pkg_font, fill=pkg_fill, align=center)

    cell(ws, row, 7, folder, font=note_font)
    cell(ws, row, 8, notes)

# Status validation
dv_status = DataValidation(type="list", formula1='"DONE,UPCOMING,IN PROGRESS,NOT STARTED"', allow_blank=True)
ws.add_data_validation(dv_status)
dv_status.add(f"D{r+1}:D{r+4}")

dv_yn = DataValidation(type="list", formula1='"YES,NO,PARTIAL"', allow_blank=True)
ws.add_data_validation(dv_yn)
dv_yn.add(f"E{r+1}:F{r+4}")

# ═══════════════════════════════════════════════════════════
# SECTION 2: PHASE 1 DELIVERABLES (what to package)
# ═══════════════════════════════════════════════════════════
r = 11
ws.merge_cells(f"A{r}:H{r}")
cell(ws, r, 1, "PHASE 1 DELIVERABLES — Package into: Deliverables from Quannex/Phase 1", font=section_font)

r = 12
h2 = ["Deliverable", "Source Material", "Source Location", "Format", "Status", "Action Needed", "Priority", "Notes"]
for c, h in enumerate(h2, 1):
    cl = ws.cell(row=r, column=c, value=h)
    cl.font = subheader_font
    cl.fill = subheader_fill
    cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cl.border = thin_border

p1_items = [
    ("Interview Synthesis Report",
     "CEN_Phase1_Interview_Synthesis.md (57 KB)",
     "Meetings & Planning/CEN_Phase1/",
     "PDF (branded)", "EXISTS — needs formatting",
     "Convert to branded PDF, add exec summary", "HIGH",
     "Core research artifact — shows CEN their organizational portrait from Phase 1"),
    ("Preliminary Face Mapping",
     "CEN_Preliminary_Face_Mapping.md (33 KB)",
     "Meetings & Planning/CEN_Phase1/",
     "PDF (branded)", "EXISTS — needs formatting",
     "Convert to branded PDF with visual dodecahedron diagram", "HIGH",
     "Shows how CEN's domains map to the 12 faces"),
    ("Meeting Transcript",
     "20-03-2026_PHASE1_Meeting-TRANSCRIPT.md (183 KB)",
     "Meetings & Planning/CEN_Phase1/",
     "PDF or MD", "EXISTS — decide if CEN-facing",
     "Review for sensitivity, redact if needed", "MEDIUM",
     "May not need to be delivered — check partnership agreement"),
    ("Personal Notes",
     "Personal Notes - Fresh, right after the meeting!.md",
     "Meetings & Planning/CEN_Phase1/",
     "N/A", "INTERNAL ONLY",
     "Do NOT deliver — keep as research evidence", "—",
     "Raw observations, thesis evidence only"),
]

for i, (deliv, source, loc, fmt, status, action, priority, notes) in enumerate(p1_items):
    row = r + 1 + i
    cell(ws, row, 1, deliv, font=bold_font)
    cell(ws, row, 2, source)
    cell(ws, row, 3, loc, font=note_font)
    cell(ws, row, 4, fmt, align=center)
    st_fill = warn_fill if "needs" in status else (done_fill if "DONE" in status else future_fill)
    cell(ws, row, 5, status, fill=st_fill)
    cell(ws, row, 6, action)
    pri_font = urgent_font if priority == "HIGH" else (warn_font if priority == "MEDIUM" else body_font)
    cell(ws, row, 7, priority, font=pri_font, align=center)
    cell(ws, row, 8, notes)

# ═══════════════════════════════════════════════════════════
# SECTION 3: PHASE 2 DELIVERABLES
# ═══════════════════════════════════════════════════════════
r = 18
ws.merge_cells(f"A{r}:H{r}")
cell(ws, r, 1, "PHASE 2 DELIVERABLES — Package into: Deliverables from Quannex/Phase 2", font=section_font)

r = 19
for c, h in enumerate(h2, 1):
    cl = ws.cell(row=r, column=c, value=h)
    cl.font = subheader_font
    cl.fill = subheader_fill
    cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cl.border = thin_border

p2_items = [
    ("Coherence Portrait / Vector Diagnostics",
     "CEN_Phase2_C1_VectorDiagnostics.md (20 KB)",
     "Meetings & Planning/CEN_Phase2/EvidencePackage/",
     "PDF (branded)", "EXISTS — needs formatting",
     "Convert to branded CEN-facing report, add visualizations", "HIGH",
     "The core deliverable — shows CEN their coherence state"),
    ("Breath Axis Analysis",
     "CEN_Phase2_C2_BreathAxes.md (21 KB)",
     "Meetings & Planning/CEN_Phase2/EvidencePackage/",
     "PDF (branded)", "EXISTS — needs formatting",
     "Format as section of coherence report or standalone", "HIGH",
     "Polarity tensions — key insight for leadership"),
    ("Divergence Narratives (Hidden Tensions)",
     "CEN_Phase2_C3_DivergenceNarratives.md (13 KB)",
     "Meetings & Planning/CEN_Phase2/EvidencePackage/",
     "PDF (branded)", "EXISTS — needs formatting",
     "Sensitive content — frame constructively for CEN", "HIGH",
     "Where co-founders diverge — handle with care"),
    ("BSC Comparison Report",
     "CEN_Phase2_C4_BSCComparison.md (27 KB)",
     "Meetings & Planning/CEN_Phase2/EvidencePackage/",
     "PDF (branded)", "EXISTS — needs formatting",
     "Format as comparison report: Spiral vs BSC (5/6 verdict)", "HIGH",
     "Academic evidence + CEN value prop: shows what Spiral adds over BSC"),
    ("Raw Scoring Data (Frozen)",
     "CEN_Phase2_RawScores_Frozen.md (9 KB)",
     "Meetings & Planning/CEN_Phase2/EvidencePackage/",
     "PDF or Excel", "EXISTS — decide if CEN-facing",
     "May include as appendix to coherence portrait", "MEDIUM",
     "Both co-founders' scores, independence verified"),
    ("Scoring Questionnaires (Completed)",
     "Dominique + Esther completed questionnaires",
     "CEN Inputs/CEN_Scoring_Questionnaire/",
     "DOCX (as-is)", "COMPLETE — already received",
     "Archive copies in Phase 2 folder for completeness", "LOW",
     "Input FROM CEN, not deliverable TO CEN — but archive"),
    ("Face Guide Companion",
     "CEN_Face_Guide_For_Scoring.pdf (202 KB)",
     "Meetings & Planning/CEN_Phase2/",
     "PDF", "COMPLETE",
     "Already formatted — copy to Phase 2 folder", "LOW",
     "Sent to CEN for scoring context"),
]

for i, (deliv, source, loc, fmt, status, action, priority, notes) in enumerate(p2_items):
    row = r + 1 + i
    cell(ws, row, 1, deliv, font=bold_font)
    cell(ws, row, 2, source)
    cell(ws, row, 3, loc, font=note_font)
    cell(ws, row, 4, fmt, align=center)
    st_fill = warn_fill if "needs" in status else (done_fill if "COMPLETE" in status else future_fill)
    cell(ws, row, 5, status, fill=st_fill)
    cell(ws, row, 6, action)
    pri_font = urgent_font if priority == "HIGH" else (warn_font if priority == "MEDIUM" else body_font)
    cell(ws, row, 7, priority, font=pri_font, align=center)
    cell(ws, row, 8, notes)

# ═══════════════════════════════════════════════════════════
# SECTION 4: PHASE 3 DELIVERABLES (upcoming)
# ═══════════════════════════════════════════════════════════
r = 28
ws.merge_cells(f"A{r}:H{r}")
cell(ws, r, 1, "PHASE 3 DELIVERABLES — Validation Session (~April 17) — Package into: Deliverables from Quannex/Phase 3", font=section_font)

r = 29
for c, h in enumerate(h2, 1):
    cl = ws.cell(row=r, column=c, value=h)
    cl.font = subheader_font
    cl.fill = subheader_fill
    cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cl.border = thin_border

p3_items = [
    ("Validation Presentation",
     "To be created from Phase 2 evidence",
     "—",
     "PPTX (branded)", "NOT STARTED",
     "Create presentation summarizing coherence portrait + key findings", "CRITICAL",
     "This is what you present to D + E at the validation meeting"),
    ("CEN-Facing Coherence Summary",
     "Synthesize from C1-C4 evidence docs",
     "—",
     "PDF (branded)", "NOT STARTED",
     "One-page or two-page executive summary for CEN leadership", "CRITICAL",
     "Leave-behind document after validation session"),
    ("Validation Session Notes",
     "Will be captured during meeting",
     "—",
     "MD then PDF", "FUTURE",
     "Record CEN's reactions, corrections, confirmations", "HIGH",
     "Critical thesis evidence — did CEN validate the portrait?"),
    ("POC Live Demo (optional)",
     "POC codebase — load CEN data into engine",
     "POC/companies/ (needs CEN template)",
     "Live demo", "REQUIRES CEN COMPANY TEMPLATE",
     "Create companies/cen/ template from scoring data", "MEDIUM",
     "Would powerfully demonstrate the tool — consider creating CEN template from Phase 2 scores"),
]

for i, (deliv, source, loc, fmt, status, action, priority, notes) in enumerate(p3_items):
    row = r + 1 + i
    cell(ws, row, 1, deliv, font=bold_font)
    cell(ws, row, 2, source)
    cell(ws, row, 3, loc, font=note_font)
    cell(ws, row, 4, fmt, align=center)
    st_fill = fail_fill if "CRITICAL" in priority else (warn_fill if "NOT STARTED" in status else future_fill)
    cell(ws, row, 5, status, fill=st_fill)
    cell(ws, row, 6, action)
    pri_font = urgent_font if priority == "CRITICAL" else (warn_font if priority in ("HIGH", "MEDIUM") else body_font)
    cell(ws, row, 7, priority, font=pri_font, align=center)
    cell(ws, row, 8, notes)

# ═══════════════════════════════════════════════════════════
# SECTION 5: PHASE 4 DELIVERABLES (future)
# ═══════════════════════════════════════════════════════════
r = 35
ws.merge_cells(f"A{r}:H{r}")
cell(ws, r, 1, "PHASE 4 DELIVERABLES — Final Report (May) — Package into: Deliverables from Quannex/Phase 4", font=section_font)

r = 36
for c, h in enumerate(h2, 1):
    cl = ws.cell(row=r, column=c, value=h)
    cl.font = subheader_font
    cl.fill = subheader_fill
    cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cl.border = thin_border

p4_items = [
    ("Final Coherence Report",
     "Full synthesis of Phase 1-3 findings",
     "—",
     "PDF (branded)", "NOT STARTED",
     "Comprehensive CEN coherence assessment with recommendations", "HIGH",
     "The capstone deliverable — CEN's return on participation"),
    ("Recommendations Document",
     "Derived from cross-case analysis + validation",
     "—",
     "PDF (branded)", "NOT STARTED",
     "Actionable recommendations based on coherence portrait", "HIGH",
     "Practical value for CEN — not just academic"),
    ("Thank You + Next Steps",
     "—", "—",
     "PDF or letter", "NOT STARTED",
     "Formal closure of research partnership", "MEDIUM",
     "Professional courtesy — partnership confirmation referenced deliverables"),
]

for i, (deliv, source, loc, fmt, status, action, priority, notes) in enumerate(p4_items):
    row = r + 1 + i
    cell(ws, row, 1, deliv, font=bold_font)
    cell(ws, row, 2, source)
    cell(ws, row, 3, loc, font=note_font)
    cell(ws, row, 4, fmt, align=center)
    cell(ws, row, 5, status, fill=future_fill)
    cell(ws, row, 6, action)
    pri_font = warn_font if priority == "HIGH" else body_font
    cell(ws, row, 7, priority, font=pri_font, align=center)
    cell(ws, row, 8, notes)

# ═══════════════════════════════════════════════════════════
# SECTION 6: CEN INPUT MATERIALS INVENTORY
# ═══════════════════════════════════════════════════════════
r = 41
ws.merge_cells(f"A{r}:H{r}")
cell(ws, r, 1, "CEN INPUT MATERIALS — Documents received from CEN (for reference)", font=section_font)

r = 42
h3 = ["Category", "Document", "Size", "Location", "Relevance", "Used In", "Notes", ""]
for c, h in enumerate(h3, 1):
    cl = ws.cell(row=r, column=c, value=h)
    cl.font = subheader_font
    cl.fill = subheader_fill
    cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cl.border = thin_border

inputs = [
    ("Legal", "Signed NDA", "461 KB", "Meeting 1st/", "Foundation", "All phases", ""),
    ("Legal", "Signed Partnership Confirmation", "409 KB", "Meeting 1st/", "Foundation", "All phases", ""),
    ("Strategy", "CEN Business Plan V2.0", "~19 MB", "CEN Inputs/Strategy/", "HIGH", "Phase 1 synthesis", "Core strategic context"),
    ("Strategy", "CEN Strategy V0.4", "425 KB", "CEN Inputs/Strategy/", "HIGH", "Phase 1 synthesis", ""),
    ("Strategy", "Purpose Vision Mission Values V0.1", "25 KB", "CEN Inputs/Vision Mission/", "HIGH", "Face mapping", ""),
    ("Pillar 1", "Conscious Leadership Course Proposal", "1.2 MB", "CEN Inputs/Pillar 1/", "MEDIUM", "Face 6 (Consciousness)", ""),
    ("Pillar 2", "AI Consulting + Governance proposals", "~230 KB", "CEN Inputs/Pillar 2/", "MEDIUM", "Face 3, Face 11", "5 documents"),
    ("Pillar 3", "LLM, DCS, Business Plans, Partners", "~42 MB", "CEN Inputs/Pillar 3/", "MEDIUM", "Multiple faces", "18 documents"),
    ("Teams", "Volunteer programs + contracts", "~470 KB", "CEN Inputs/Teams/", "MEDIUM", "Face 4, Face 9", "5 documents"),
    ("Scoring", "Dominique questionnaire (completed)", "27 KB", "CEN Inputs/Scoring/", "CRITICAL", "Phase 2 evidence", "Returned Apr 7"),
    ("Scoring", "Esther questionnaire (completed)", "30 KB", "CEN Inputs/Scoring/", "CRITICAL", "Phase 2 evidence", "Returned Apr 7"),
    ("Recording", "Phase 1 Interview (video)", "4.08 GB", "CEN_Phase1/", "CRITICAL", "Phase 1 synthesis", "Primary research data"),
    ("Recording", "Phase 2 Meeting (video)", "430 MB", "CEN_Phase2/", "CRITICAL", "Phase 2 evidence", "Primary research data"),
]

for i, (cat, doc, size, loc, rel, used, notes) in enumerate(inputs):
    row = r + 1 + i
    cell(ws, row, 1, cat)
    cell(ws, row, 2, doc, font=bold_font)
    cell(ws, row, 3, size, align=center)
    cell(ws, row, 4, loc, font=note_font)
    rel_font = urgent_font if rel == "CRITICAL" else (warn_font if rel == "HIGH" else body_font)
    cell(ws, row, 5, rel, font=rel_font, align=center)
    cell(ws, row, 6, used)
    cell(ws, row, 7, notes)

# ── Summary at bottom ──
r = 57
ws.merge_cells(f"A{r}:H{r}")
cell(ws, r, 1,
     "KEY GAP: Phase 1 and Phase 2 work is DONE but delivery folders are EMPTY. "
     "Package existing evidence into branded PDFs before Phase 3 validation (~April 17).",
     font=urgent_font)

r = 59
ws.merge_cells(f"A{r}:H{r}")
cell(ws, r, 1,
     "Folder path: Final Thesis/Thesis Work/External Partners/Coherence Partnership Quannex x CEN/"
     "Local Source for CEN facing documents/CEN Facing Folder/Deliverables from Quannex/",
     font=note_font)

# ── Column widths ──
col_widths = [30, 35, 14, 16, 22, 38, 12, 45]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A5"

# ── Also update Dashboard Sheet 1 with corrected milestones ──
ws1 = wb["Dashboard"]
# Fix the milestone line — CEN Phase 2 is done, next is Phase 3
ws1["B7"].value = "Phase 3 validation — ~April 17"
ws1["B7"].font = Font(name="Segoe UI", size=10, color=DARK)

# ── Also update Thesis Evidence sheet milestones ──
ws4 = wb["Thesis Evidence"]
# Find the CEN Phase 2 timeline row and update it
for row in range(1, ws4.max_row + 1):
    val = ws4.cell(row=row, column=1).value
    if val and "April 14" in str(val):
        ws4.cell(row=row, column=1).value = "April 14 (DONE)"
        ws4.cell(row=row, column=2).value = "CEN Phase 2 scoring — COMPLETE (session 25, Apr 7)"
        ws4.cell(row=row, column=3).value = "BSC comparison 5/6. Evidence package ready."
        break

wb.save(path)
print(f"Sheet 'CEN Deliverables' added to CONTROL_PANEL.xlsx")
print(f"Dashboard milestone updated: Phase 3 validation ~April 17")
print(f"Thesis Evidence timeline corrected: CEN Phase 2 marked DONE")
