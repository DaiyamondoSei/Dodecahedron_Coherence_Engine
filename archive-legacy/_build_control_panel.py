"""
Build CONTROL_PANEL.xlsx — the POC nerve center.
12 faces × 5 elements = 60 dimensions of quality.
Run once, then delete this script.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import date, datetime

wb = openpyxl.Workbook()

# ── Brand colors ──
CYAN = "00BCD4"
DARK_CYAN = "00838F"
LIGHT_CYAN = "E0F7FA"
GREEN = "4CAF50"
YELLOW = "FFC107"
RED = "F44336"
GRAY = "9E9E9E"
WHITE = "FFFFFF"
DARK = "263238"

# ── Reusable styles ──
header_font = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
header_fill = PatternFill(start_color=DARK_CYAN, end_color=DARK_CYAN, fill_type="solid")
subheader_fill = PatternFill(start_color=CYAN, end_color=CYAN, fill_type="solid")
subheader_font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
body_font = Font(name="Segoe UI", size=10, color=DARK)
title_font = Font(name="Segoe UI", size=16, bold=True, color=DARK_CYAN)
subtitle_font = Font(name="Segoe UI", size=12, bold=True, color=DARK)
note_font = Font(name="Segoe UI", size=9, italic=True, color=GRAY)
thin_border = Border(
    left=Side(style="thin", color="BDBDBD"),
    right=Side(style="thin", color="BDBDBD"),
    top=Side(style="thin", color="BDBDBD"),
    bottom=Side(style="thin", color="BDBDBD"),
)
wrap_align = Alignment(wrap_text=True, vertical="top")
center_align = Alignment(horizontal="center", vertical="center")

# Status fill colors
pass_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
warn_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")
fail_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
unchecked_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")


def style_header_row(ws, row, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border


def style_body_cell(ws, row, col, value=None):
    cell = ws.cell(row=row, column=col)
    if value is not None:
        cell.value = value
    cell.font = body_font
    cell.alignment = wrap_align
    cell.border = thin_border
    return cell


def add_status_validation(ws, cell_range):
    dv = DataValidation(
        type="list",
        formula1='"PASS,WARN,FAIL,N/A,—"',
        allow_blank=True,
    )
    dv.error = "Please select PASS, WARN, FAIL, N/A, or —"
    dv.errorTitle = "Invalid Status"
    ws.add_data_validation(dv)
    dv.add(cell_range)


# ═══════════════════════════════════════════════════════════
# SHEET 1: DASHBOARD
# ═══════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Dashboard"
ws1.sheet_properties.tabColor = CYAN

# Title
ws1.merge_cells("A1:G1")
ws1["A1"].value = "QUANNEX POC — CONTROL PANEL"
ws1["A1"].font = title_font
ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")

# Thesis countdown
ws1.merge_cells("A3:B3")
ws1["A3"].value = "THESIS DEFENSE COUNTDOWN"
ws1["A3"].font = subtitle_font

deadline = date(2026, 6, 1)
today = date.today()
days_left = (deadline - today).days
week_num = 15 - (days_left // 7)

countdown_data = [
    ("Days Remaining", str(days_left)),
    ("Current Week", f"{week_num} of 15"),
    ("Hard Deadline", "June 1, 2026 (noon)"),
    ("Next Milestone", "CEN Phase 2 — April 14"),
]
for i, (label, val) in enumerate(countdown_data):
    r = 4 + i
    ws1.cell(row=r, column=1, value=label).font = Font(name="Segoe UI", size=10, bold=True, color=DARK)
    c = ws1.cell(row=r, column=2, value=val)
    c.font = body_font
    if "Days" in label:
        c.font = Font(name="Segoe UI", size=14, bold=True, color=RED if days_left < 30 else DARK_CYAN)

# 12 Faces table
face_start = 9
headers = ["Face #", "Name", "Dimension", "Status", "Key Metric", "Last Assessed", "Notes"]
for col, h in enumerate(headers, 1):
    ws1.cell(row=face_start, column=col, value=h)
style_header_row(ws1, face_start, len(headers))

faces = [
    (1, "Survival", "Correctness", "PASS", "0 syntax errors", "2026-04-11", ""),
    (2, "Foundation", "Structural Integrity", "PASS", "PhiHarmonics consistent", "2026-04-11", ""),
    (3, "Clarity", "Documentation", "WARN", "α-drift class patched", "2026-04-07", "Doc-drift risk ongoing"),
    (4, "Harmony", "DRY & Hygiene", "PASS", "Logger migration done", "2026-03-14", "3 spectral files = tech debt"),
    (5, "Signal", "Honest Output", "PASS", "Test honesty verified", "2026-02-08", "Stress test was once dishonest"),
    (6, "Consciousness", "Intelligence Interop", "PASS", "CLAUDE.md + hooks active", "2026-04-07", ""),
    (7, "Coherence", "Pattern Alignment", "PASS", "5-pass engine intact", "2026-04-07", ""),
    (8, "Resilience", "Risk & Security", "WARN", "Circuit breaker active", "2026-02-08", "Edge cases need re-testing"),
    (9, "Growth", "Regenerative Quality", "PASS", "Self-dev pattern installed", "2026-04-07", "13 growth ideas pending"),
    (10, "Integrity", "Test Coverage", "PASS", "306/306 green", "2026-04-11", "Smoke test needs HTTP server"),
    (11, "Discovery", "Research & Creativity", "PASS", "11 emergent truths found", "2026-03-14", ""),
    (12, "Radiance", "Sustainable Exponentiality", "PASS", "Spiral loop self-applying", "2026-04-07", "Law 1 inscribed"),
]

for i, (fnum, name, dim, status, metric, assessed, notes) in enumerate(faces):
    r = face_start + 1 + i
    vals = [fnum, name, dim, status, metric, assessed, notes]
    for col, v in enumerate(vals, 1):
        cell = style_body_cell(ws1, r, col, v)
        if col == 1:
            cell.alignment = center_align
        if col == 4:  # status column
            if v == "PASS":
                cell.fill = pass_fill
            elif v == "WARN":
                cell.fill = warn_fill
            elif v == "FAIL":
                cell.fill = fail_fill

# Status validation for column D
add_status_validation(ws1, f"D{face_start+1}:D{face_start+12}")

# Overall verdict
verdict_row = face_start + 14
ws1.merge_cells(f"A{verdict_row}:C{verdict_row}")
ws1.cell(row=verdict_row, column=1, value="OVERALL VERDICT").font = subtitle_font
ws1.cell(row=verdict_row, column=4, value="CLEAN").font = Font(
    name="Segoe UI", size=14, bold=True, color="00838F"
)

# Law 1 reminder
law_row = verdict_row + 2
ws1.merge_cells(f"A{law_row}:G{law_row}")
ws1[f"A{law_row}"].value = (
    'Law 1: "Substrate failures are invisible until asked about. '
    'Therefore: ask about them on a schedule, not on suspicion."'
)
ws1[f"A{law_row}"].font = note_font

# Column widths
col_widths = [8, 18, 25, 10, 28, 15, 30]
for i, w in enumerate(col_widths, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

ws1.freeze_panes = f"A{face_start+1}"

# ═══════════════════════════════════════════════════════════
# SHEET 2: 60 ELEMENTS
# ═══════════════════════════════════════════════════════════
ws2 = wb.create_sheet("60 Elements")
ws2.sheet_properties.tabColor = "7C4DFF"

elements_per_face = ["Earth", "Water", "Fire", "Air", "Ether"]

# All 60 element descriptions (from SKILL.md)
element_checks = {
    1: [
        "Syntax validation — node --check on every changed .js file",
        "No crashes — POC loads in browser without console errors",
        "Script & dependency order — spectral/breath analyzers before main.js",
        "Environment compatibility — Windows path separators, pwsh vs bash",
        "Import resolution — ES module imports, cache-busting ?v= strings",
    ],
    2: [
        "Variable shadowing — locals shadowing outer scope (caused real bugs)",
        "PhiHarmonics constants — no hardcoded 0.618, 0.382, 1.618 etc.",
        "Export contracts — window globals: Quannex, PhiHarmonics, Logger etc.",
        "Type safety — face IDs numbers (1-12), energy 0-1, NaN guards",
        "Import hygiene — unused imports removed, no side-effect-only imports",
    ],
    3: [
        '"Why" comments — Greek params (α-θ) documented, phi powers explained',
        "Architecture docs — FILE_STRUCTURE_MAP, MODULE_ARCHITECTURE current",
        "Debug trails — symptom → root cause → fix documented",
        "Calculation audit trail — every formula traceable input→formula→output",
        "Cross-workspace notes — thesis-relevant evidence flagged",
    ],
    4: [
        "Deduplication — 3 spectral analyzer files consistent",
        "Dead code removal — no commented-out code, use archive/ dir",
        "Naming conventions — camelCase vars, PascalCase classes, kebab files",
        "Anti-pattern cleanup — Logger.* not console.log, null checks on DOM",
        "File organization — JS in correct subdirectory, pages in pages/",
    ],
    5: [
        "Test honesty — no try/catch swallowing failures in tests",
        "Error message accuracy — Logger.warn/error describe actual problem",
        "Logging consistency — all through Logger.*, correct severity levels",
        "Data validation — data-validator catches corruption patterns",
        "No silent failures — edge cases produce warnings, not wrong answers",
    ],
    6: [
        "AI navigability — new Claude session orients in <30 seconds",
        "Human readability — phi ratios explained, metaphor accessible",
        "Dependency maps — loader→data→engine→analysis chain discoverable",
        "Entry points marked — orchestrator, demo, pages, engine API documented",
        "Session context — hooks provide context, memory MCP populated",
    ],
    7: [
        "Engine patterns — follows 5-pass calculation pattern in main.js",
        "Constant usage — PhiHarmonics for thresholds, TuningConfig for params",
        "Analyzer structure — class with analyze() method, structured results",
        "Sacred geometry respect — faces/edges/vertices vocabulary, not generic",
        "Data format consistency — company templates follow quannex/ format",
    ],
    8: [
        "Edge cases — 0-energy faces, missing KPI data, NaN propagation",
        "Stress testing — all 4 templates, rapid updateKPI() calls",
        "Security — CSV injection safe, sessionStorage validated, no path traversal",
        "Failure mode simulation — missing files, malformed JSON, full storage",
        "Risk register — enumerate all risks, severity, mitigation path",
    ],
    9: [
        "Tech debt awareness — new debt created? Known: 3 spectral files, CSV legacy",
        "Extensibility — new templates without code changes? New faces beyond 12?",
        "Proactive improvements — clarified messages, tightened null checks?",
        "TODO tracking — all TODOs in report, not left as // TODO comments",
        "Sustainability — main.js (1917 lines) appropriate, or time to extract?",
    ],
    10: [
        "Suite results — Phi:127 Integration:55 Excel:42 Validator:82 Smoke:77",
        "New function coverage — every new public function has a test",
        "Regression prevention — tests against all 4 company templates",
        "Test quality — math tests verify identities, not just 'returns a number'",
        "Evidence gathering — thesis-relevant test results captured",
    ],
    11: [
        "Internet research — searched for existing solutions to hard problems?",
        "Out-of-box thinking — simpler formulation? More elegant approach?",
        "Helicopter view — serves bigger picture? Thesis demonstration?",
        "Agent team usage — complex changes used research/impl/review agents?",
        "Cross-workspace pollination — benefits Thesis or Business Data?",
    ],
    12: [
        "Compounding value — does this make NEXT Quannex work easier?",
        "Capability building — reusable utilities extracted to js/core/?",
        "Pipeline development — new tests, skills, hooks, templates?",
        "Highest-leverage action — ONE thing with greatest compounding effect?",
        "Future acceleration — next session feels welcomed and empowered?",
    ],
}

face_names = {
    1: "Survival", 2: "Foundation", 3: "Clarity", 4: "Harmony",
    5: "Signal", 6: "Consciousness", 7: "Coherence", 8: "Resilience",
    9: "Growth", 10: "Integrity", 11: "Discovery", 12: "Radiance",
}

headers2 = ["Face #", "Face Name", "Element", "Check Description", "Status", "Evidence / Notes", "Last Checked"]
for col, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=col, value=h)
style_header_row(ws2, 1, len(headers2))

row = 2
for face_num in range(1, 13):
    for elem_idx, elem_name in enumerate(elements_per_face):
        style_body_cell(ws2, row, 1, face_num).alignment = center_align
        style_body_cell(ws2, row, 2, face_names[face_num])
        style_body_cell(ws2, row, 3, elem_name)
        style_body_cell(ws2, row, 4, element_checks[face_num][elem_idx])
        style_body_cell(ws2, row, 5, "—")
        style_body_cell(ws2, row, 6, "")
        style_body_cell(ws2, row, 7, "")
        # Light row shading to separate faces
        if face_num % 2 == 0:
            for c in range(1, 8):
                ws2.cell(row=row, column=c).fill = PatternFill(
                    start_color="FAFAFA", end_color="FAFAFA", fill_type="solid"
                )
        row += 1

add_status_validation(ws2, f"E2:E{row-1}")

col_widths2 = [8, 15, 10, 55, 10, 30, 14]
for i, w in enumerate(col_widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"

# ═══════════════════════════════════════════════════════════
# SHEET 3: TEST TRACKER
# ═══════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Test Tracker")
ws3.sheet_properties.tabColor = GREEN

headers3 = ["Suite", "File Path", "Expected", "Actual", "Pass", "Fail", "Last Run", "Notes"]
for col, h in enumerate(headers3, 1):
    ws3.cell(row=1, column=col, value=h)
style_header_row(ws3, 1, len(headers3))

test_data = [
    ("Phi Math Constants", "tests/phi-math.test.js", 127, 127, 127, 0, "2026-04-11", ""),
    ("Integration (Engine)", "tests/integration.test.mjs", 55, 55, 55, 0, "2026-04-11", "MODULE_TYPELESS warning (cosmetic)"),
    ("Excel Parser Round-Trip", "tests/excel-parser.test.mjs", 42, 42, 42, 0, "2026-04-11", ""),
    ("Data Validator", "tests/data-validator.test.mjs", 82, 82, 82, 0, "2026-04-11", "Was 88 in some old docs — 82 is the real count"),
    ("Browser Smoke Test", "tests/smoke-test.mjs", 77, "—", "—", "—", "—", "Needs HTTP server on localhost:8000"),
]

for i, row_data in enumerate(test_data):
    r = 2 + i
    for col, v in enumerate(row_data, 1):
        cell = style_body_cell(ws3, r, col, v)
        if col in (3, 4, 5, 6):
            cell.alignment = center_align

# Totals row
r_total = 8
ws3.cell(row=r_total, column=1, value="TOTAL").font = Font(name="Segoe UI", size=10, bold=True, color=DARK)
style_body_cell(ws3, r_total, 2, "")
style_body_cell(ws3, r_total, 3, 383).alignment = center_align
style_body_cell(ws3, r_total, 4, 306).alignment = center_align
style_body_cell(ws3, r_total, 5, 306).alignment = center_align
style_body_cell(ws3, r_total, 6, 0).alignment = center_align
ws3.cell(row=r_total, column=3).font = Font(name="Segoe UI", size=10, bold=True, color=DARK)
ws3.cell(row=r_total, column=4).font = Font(name="Segoe UI", size=10, bold=True, color=DARK)
ws3.cell(row=r_total, column=5).font = Font(name="Segoe UI", size=10, bold=True, color="00838F")
ws3.cell(row=r_total, column=6).font = Font(name="Segoe UI", size=10, bold=True, color=DARK)
style_body_cell(ws3, r_total, 7, "2026-04-11")
style_body_cell(ws3, r_total, 8, "Smoke test not included (needs server)")

# HTML/Legacy tests section
r_html = 10
ws3.merge_cells(f"A{r_html}:H{r_html}")
ws3[f"A{r_html}"].value = "HTML Browser Tests (manual, open in browser)"
ws3[f"A{r_html}"].font = subheader_font
ws3[f"A{r_html}"].fill = subheader_fill

html_tests = [
    ("Circuit Breaker", "tests/circuit-breaker.test.html", "—", "—", "—", "—", "—", "Needs browser"),
    ("KPI Constants Integrity", "tests/kpi-constants-integrity.test.html", "—", "—", "—", "—", "—", "Needs browser"),
    ("Phase 3 Integrity", "tests/phase3-integrity-tests.html", "—", "—", "—", "—", "—", "Needs browser"),
]
for i, row_data in enumerate(html_tests):
    r = r_html + 1 + i
    for col, v in enumerate(row_data, 1):
        cell = style_body_cell(ws3, r, col, v)
        if col in (3, 4, 5, 6):
            cell.alignment = center_align

# How to run
r_how = r_html + 5
ws3.merge_cells(f"A{r_how}:H{r_how}")
ws3[f"A{r_how}"].value = "How to run:  node tests/run-all.js    (for smoke test: first run  python -m http.server 8000)"
ws3[f"A{r_how}"].font = note_font

col_widths3 = [25, 30, 10, 10, 10, 8, 14, 40]
for i, w in enumerate(col_widths3, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.freeze_panes = "A2"

# ═══════════════════════════════════════════════════════════
# SHEET 4: THESIS EVIDENCE MAP
# ═══════════════════════════════════════════════════════════
ws4 = wb.create_sheet("Thesis Evidence")
ws4.sheet_properties.tabColor = "FF6F00"

# Title
ws4.merge_cells("A1:G1")
ws4["A1"].value = "POC → THESIS EVIDENCE MAP"
ws4["A1"].font = title_font
ws4["A1"].alignment = Alignment(horizontal="center")

headers4 = ["POC Artifact", "File Path", "Size", "Status", "Target Chapter", "Integration Status", "Notes"]
for col, h in enumerate(headers4, 1):
    ws4.cell(row=3, column=col, value=h)
style_header_row(ws4, 3, len(headers4))

thesis_data = [
    ("Spectral Analysis Chapter", "docs/thesis/SPECTRAL_ANALYSIS_CHAPTER.md", "36 KB", "READY", "Ch4/5 Results", "Not yet integrated", "Formatted for .docx import"),
    ("Consciousness Model", "docs/thesis/CONSCIOUSNESS_MODEL.md", "12 KB", "READY", "Ch2/3 Literature", "Not yet integrated", "7 octaves + psychology basis"),
    ("Defense Preparation", "docs/thesis/DEFENSE_PREPARATION.md", "30 KB", "READY", "Defense prep (study guide)", "N/A", "Committee Q&A, 499 lines"),
    ("Emergent Mathematical Truths", "docs/thesis/EMERGENT_MATHEMATICAL_TRUTHS.md", "50 KB", "READY", "Ch4 Discussion", "Not yet integrated", "11 discoveries, stress test evidence"),
    ("Novel Mathematical Contributions", "docs/thesis/NOVEL_MATHEMATICAL_CONTRIBUTIONS.md", "36 KB", "READY", "Ch4 Appendix", "Not yet integrated", "4 frameworks, ALCOA+ rigor"),
    ("Wisdom Brief", "docs/thesis/WISDOM_BRIEF.md", "17 KB", "READY", "Ch2 Literature", "Not yet integrated", "Philosophical grounding"),
    ("Calculation Audit Trail", "docs/math/CALCULATION_AUDIT_TRAIL.md", "28 KB", "READY", "Appendix", "Not yet integrated", '"Academic armor" — every formula traced'),
    ("Sacred Geometry Proof", "docs/math/SACRED_GEOMETRY_PROOF.md", "31 KB", "READY", "Ch2/Appendix", "Not yet integrated", "Eigenvalue verification, phi proofs"),
    ("Spectral Implementation", "docs/math/SPECTRAL_IMPLEMENTATION.md", "8 KB", "READY", "Ch4 Methods", "Not yet integrated", "Graph Laplacian → code mapping"),
    ("Thesis Audit Report", "THESIS_AUDIT_REPORT.md", "25 KB", "READY", "Reference only", "N/A", "VERDICT: PRISTINE, zero findings"),
]

for i, row_data in enumerate(thesis_data):
    r = 4 + i
    for col, v in enumerate(row_data, 1):
        cell = style_body_cell(ws4, r, col, v)
        if col == 4:
            cell.fill = pass_fill

# Integration status validation
int_dv = DataValidation(
    type="list",
    formula1='"Not yet integrated,In progress,Integrated,N/A"',
    allow_blank=True,
)
ws4.add_data_validation(int_dv)
int_dv.add(f"F4:F{4+len(thesis_data)-1}")

# Timeline section
r_tl = 4 + len(thesis_data) + 2
ws4.merge_cells(f"A{r_tl}:G{r_tl}")
ws4[f"A{r_tl}"].value = "CRITICAL TIMELINE"
ws4[f"A{r_tl}"].font = subtitle_font

timeline = [
    ("April 14, 2026", "CEN Phase 2 scoring deadline", "BLOCKING — no slip tolerance"),
    ("April 15-30", "Validation session with CEN", ""),
    ("May 1-20", "Write Chapters 5-7 (Case Study B, Cross-Case, Conclusions)", "PRIMARY WRITING WINDOW"),
    ("May 25", "Final review pass", ""),
    ("June 1 (noon)", "HARD DEADLINE — upload to Onstage", "NON-NEGOTIABLE"),
]

tl_headers = ["Date", "Milestone", "Notes"]
for col, h in enumerate(tl_headers, 1):
    ws4.cell(row=r_tl + 1, column=col, value=h)
    ws4.cell(row=r_tl + 1, column=col).font = subheader_font
    ws4.cell(row=r_tl + 1, column=col).fill = subheader_fill

for i, (d, m, n) in enumerate(timeline):
    r = r_tl + 2 + i
    style_body_cell(ws4, r, 1, d)
    style_body_cell(ws4, r, 2, m)
    style_body_cell(ws4, r, 3, n)
    if "HARD" in d or "BLOCKING" in n:
        for c in range(1, 4):
            ws4.cell(row=r, column=c).font = Font(name="Segoe UI", size=10, bold=True, color="D32F2F")

# Blockers
r_bl = r_tl + 2 + len(timeline) + 1
ws4.merge_cells(f"A{r_bl}:G{r_bl}")
ws4[f"A{r_bl}"].value = "ACTIVE BLOCKERS"
ws4[f"A{r_bl}"].font = Font(name="Segoe UI", size=12, bold=True, color="D32F2F")

blockers = [
    ("IFC language reframe (Ch2-5)", "Blocked on Jan Willem confirmation", ""),
    ("Breath Axes scoring + POC run", "Needed for Chapter 5 §5.1.3", ""),
]
for i, (b, detail, note) in enumerate(blockers):
    r = r_bl + 1 + i
    style_body_cell(ws4, r, 1, b).font = Font(name="Segoe UI", size=10, bold=True, color=DARK)
    style_body_cell(ws4, r, 2, detail)
    style_body_cell(ws4, r, 3, note)

col_widths4 = [30, 40, 8, 10, 22, 20, 38]
for i, w in enumerate(col_widths4, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w
ws4.freeze_panes = "A4"

# ═══════════════════════════════════════════════════════════
# SHEET 5: CODEBASE INVENTORY
# ═══════════════════════════════════════════════════════════
ws5 = wb.create_sheet("Codebase Inventory")
ws5.sheet_properties.tabColor = "5C6BC0"

headers5 = ["Module", "Directory", "Files", "Key Components", "Status", "Notes"]
for col, h in enumerate(headers5, 1):
    ws5.cell(row=1, column=col, value=h)
style_header_row(ws5, 1, len(headers5))

inventory = [
    ("Core Engine", "js/core/", "~10", "Face.js, TuningConfig.js, KPI.js", "Healthy", "SSOT for Greek constants (α-θ)"),
    ("Main Engine", "js/main.js", "1", "1917-line calculation engine (5-pass)", "Healthy", "Consider extracting to js/core/"),
    ("Advanced Analysis", "js/advanced/", "~8", "spectral-analyzer, edge-analyzer, vertex-analyzer", "Healthy", "3 spectral files = unification debt"),
    ("Constants", "js/constants/", "~5", "phi-harmonics.js, sacred-inquiry.js", "Healthy", "SSOT for phi-derived values"),
    ("Data System", "js/data-system/", "~6", "data-validator.js, unified-data-loader.js", "Healthy", ""),
    ("Simulator", "js/simulator/", "9", "sim-*.js modules", "Healthy", "Gold-documented (Feb 8)"),
    ("Edge System", "js/edge/", "~4", "unified-edge.js, edge-manager.js", "Healthy", "Pure Membrane Model"),
    ("Shadow System", "js/shadow/", "14", "Shadow detection + generation", "Healthy", "AI-integrated"),
    ("Orchestrator", "js/orchestrator/", "24", "5-step wizard flow", "Healthy", "Has archive/ subdir"),
    ("AI Layer", "js/ai/", "36", "Gemini integration, mapping context", "Healthy", ""),
    ("Pages", "pages/", "9", "HTML entry points + orchestrator", "Healthy", ""),
    ("Tests", "tests/", "10", "5 suites + HTML tests", "306/306 green", "See Test Tracker sheet"),
    ("Companies", "companies/", "4 dirs", "quannex, nova-tech, apex, zenith", "Healthy", "Each: json + csv + mapping"),
    ("Documentation", "docs/", "~60", "thesis/, math/, edge/, archive/", "WARN", "13 growth ideas from Feb 16 audit"),
    ("Math Docs", "math/", "~6", "PENTAGRAM_ANALYSIS, OCTAVE_FRAMEWORK", "Healthy", "Standalone math references"),
]

for i, row_data in enumerate(inventory):
    r = 2 + i
    for col, v in enumerate(row_data, 1):
        cell = style_body_cell(ws5, r, col, v)
        if col == 3:
            cell.alignment = center_align
        if col == 5:
            if "Healthy" in str(v) or "green" in str(v):
                cell.fill = pass_fill
            elif v == "WARN":
                cell.fill = warn_fill

# Summary
r_sum = 2 + len(inventory) + 1
ws5.merge_cells(f"A{r_sum}:F{r_sum}")
ws5[f"A{r_sum}"].value = "TOTALS: 195 JS files | 9 pages | 4 company templates | ~60 docs | 306 automated tests"
ws5[f"A{r_sum}"].font = Font(name="Segoe UI", size=10, bold=True, color=DARK_CYAN)

col_widths5 = [18, 20, 8, 42, 14, 38]
for i, w in enumerate(col_widths5, 1):
    ws5.column_dimensions[get_column_letter(i)].width = w
ws5.freeze_panes = "A2"

# ═══════════════════════════════════════════════════════════
# SHEET 6: SESSION LOG
# ═══════════════════════════════════════════════════════════
ws6 = wb.create_sheet("Session Log")
ws6.sheet_properties.tabColor = "EC407A"

headers6 = ["Date", "Focus", "Key Changes", "Commits", "Tests", "Spiral Verdict", "Thesis Impact"]
for col, h in enumerate(headers6, 1):
    ws6.cell(row=1, column=col, value=h)
style_header_row(ws6, 1, len(headers6))

# Pre-populate from git log
sessions = [
    ("2026-04-11", "Control Panel", "Created CONTROL_PANEL.xlsx", "—", "306/306", "—", "Tracking artifact created"),
    ("2026-04-07", "α-drift fix + Spiral Loop", "Patched 5 doc files, installed spiral refinement", "e8e99d1", "306/306", "CLEAN", "Doc-drift class documented"),
    ("2026-03-14", "Spectral eigenvalue fix", "Corrected Laplacian topology across 3 analyzers", "7019657", "306/306", "—", "Spectral chapter updated"),
    ("2026-03-14", "Sensitivity matrix", "Multi-mode sensitivity, KPI sensitivity analysis", "6f0ccad", "306/306", "—", ""),
    ("2026-02-16", "Documentation audit", "5-agent parallel audit, 8 priority actions", "—", "306/306", "—", "Audit tracker created"),
    ("2026-02-08", "Stress test + simulator", "Full vision coherence simulator, honest stress report", "9106dbf", "306/306", "—", "Emergent truths discovered"),
    ("2026-02-01", "Thesis audit", "Complete thesis readiness audit", "—", "306/306", "—", "VERDICT: PRISTINE"),
    ("2026-01-31", "Vertex topology fix", "Unified vertex topology across codebase", "22db8cd", "306/306", "—", ""),
]

for i, row_data in enumerate(sessions):
    r = 2 + i
    for col, v in enumerate(row_data, 1):
        style_body_cell(ws6, r, col, v)

# Empty rows for future sessions
for i in range(15):
    r = 2 + len(sessions) + i
    for col in range(1, 8):
        cell = style_body_cell(ws6, r, col, "")

# Verdict validation
verdict_dv = DataValidation(
    type="list",
    formula1='"RADIANT,CLEAN,NEEDS FIXES,BLOCKED,—"',
    allow_blank=True,
)
ws6.add_data_validation(verdict_dv)
verdict_dv.add(f"F2:F{2+len(sessions)+14}")

# How to use note
r_note = 2 + len(sessions) + 16
ws6.merge_cells(f"A{r_note}:G{r_note}")
ws6[f"A{r_note}"].value = "HOW TO USE: After each session, add one row. Date + Focus + Key Changes are the minimum. Run tests and record. Spiral Verdict from /selfcritique."
ws6[f"A{r_note}"].font = note_font

col_widths6 = [14, 25, 40, 12, 12, 16, 30]
for i, w in enumerate(col_widths6, 1):
    ws6.column_dimensions[get_column_letter(i)].width = w
ws6.freeze_panes = "A2"

# ═══════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════
output_path = r"C:\Users\murau\OneDrive\Stalinis kompiuteris\POC\CONTROL_PANEL.xlsx"
wb.save(output_path)
print(f"✅ CONTROL_PANEL.xlsx created at: {output_path}")
print(f"   Sheets: {wb.sheetnames}")
print(f"   60 Elements: {row - 2} rows")
print(f"   Thesis countdown: {days_left} days to June 1, 2026")
