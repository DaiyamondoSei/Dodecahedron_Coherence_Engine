"""
Build POC_420_Songbook_Phase4_Commitment.xlsx — 7-sheet CEN-facing commitment artifact.

Usage:
    python _build_songbook_phase4_commitment.py
    python _build_songbook_phase4_commitment.py --spiral-report PATH
"""
import sys, json, re, argparse, logging
from pathlib import Path

if sys.version_info < (3, 9):
    print("[ERROR] Python 3.9+ required."); sys.exit(1)
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[ERROR] openpyxl not installed. Run: pip install openpyxl"); sys.exit(1)

logging.basicConfig(level=logging.INFO, format="[%(levelname)s] %(asctime)s %(message)s",
                    datefmt="%Y-%m-%dT%H:%M:%S")
log = logging.getLogger(__name__)

BASE     = Path(r"C:\Users\murau\OneDrive\Stalinis kompiuteris\POC")
JSON_PATH = BASE / "data" / "json" / "refrence_models_extracted.json"
OUT_PATH  = BASE / "deliverables" / "POC_420_Songbook_Phase4_Commitment.xlsx"
TMP_PATH  = OUT_PATH.with_name(OUT_PATH.stem + "_tmp.xlsx")

# Brand colors (mirrored from _build_control_panel.py)
C = dict(
    cyan="00BCD4", dcyan="00838F", lcyan="E0F7FA",
    lgreen="E8F5E9", lyellow="FFF8E1", orange="FF6F00",
    worange="FF8F00", lorange="FFF3E0", red="F44336",
    lred="FFEBEE", gray="9E9E9E", lgray="F5F5F5",
    white="FFFFFF", dark="263238", purple="7C4DFF",
)

DOMAIN = {1:"Financial Capital",2:"Intellectual Capital",3:"Human Capital",
          4:"Structural Capital",5:"Market Resonance",6:"Community & Partners",
          7:"Brand & Reputation",8:"Core Operations",9:"Regenerative Flow",
          10:"Foundational Values",11:"Funding Pipeline",12:"Risk & Resilience"}
OCTAVE = {1:"Survival",2:"Structure",3:"Relationships",4:"Creativity",
          5:"Expression",6:"Vision",7:"Radiance"}
CEN_SCORES = {1:2,2:5,3:4,4:2,5:2,6:4,7:3,8:2,9:8,10:8,11:2,12:3}
WORKED_COORDS = {(1,1,"Earth"),(6,1,"Earth"),(12,1,"Earth"),(2,3,"Air"),(10,2,"Ether")}

# Style helpers
def fl(h): return PatternFill(start_color=h, end_color=h, fill_type="solid")
def fn(sz=10, bold=False, color=None, italic=False):
    return Font(name="Segoe UI", size=sz, bold=bold, color=color or C["dark"], italic=italic)
def br():
    s = Side(style="thin", color="BDBDBD")
    return Border(left=s, right=s, top=s, bottom=s)
def al(h="left", v="top", wrap=True): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def cw(ws, widths):
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w

def hdr(ws, row, ncols, bg=None, fg=None, sz=11):
    for col in range(1, ncols+1):
        c = ws.cell(row=row, column=col)
        c.font = fn(sz=sz, bold=True, color=fg or C["white"])
        c.fill = fl(bg or C["dcyan"])
        c.alignment = al(h="center", v="center"); c.border = br()

def cell(ws, row, col, val=None, bold=False, bg=None, clr=None, h="left", sz=10):
    c = ws.cell(row=row, column=col)
    if val is not None: c.value = val
    c.font = fn(sz=sz, bold=bold, color=clr or C["dark"])
    c.alignment = al(h=h); c.border = br()
    if bg: c.fill = fl(bg)
    return c

def sec(ws, row, text, bg=None, fg=None, sz=12, ncols=6, height=20):
    ws.merge_cells(f"A{row}:{get_column_letter(ncols)}{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.font = fn(sz=sz, bold=True, color=fg or C["white"])
    c.fill = fl(bg or C["dcyan"]); c.alignment = al(h="center", v="center")
    ws.row_dimensions[row].height = height; return row+1

def blob(ws, row, text, ncols=6, bg=None, height=16, italic=False):
    ws.merge_cells(f"A{row}:{get_column_letter(ncols)}{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.font = fn(sz=10, italic=italic, color=C["dark"])
    c.fill = fl(bg or C["lcyan"]); c.alignment = al(h="left", v="top")
    c.border = br(); ws.row_dimensions[row].height = height; return row+1

# ── Sheet 1: Cover ────────────────────────────────────────────────────────────
def build_cover(wb):
    ws = wb.active; ws.title = "Cover"; ws.sheet_properties.tabColor = C["dcyan"]
    ws.merge_cells("A1:G1")
    c = ws["A1"]; c.value = "Spiral Octave Songbook \u2014 Phase 4 Commitment Artifact"
    c.font = fn(sz=18, bold=True, color=C["dcyan"]); c.alignment = al(h="center",v="center")
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A2:G2"); c = ws["A2"]
    c.value = "CEN Research-Partner Meeting  |  2026-04-17  |  Version 1.0  |  Date: 2026-04-15"
    c.font = fn(sz=11, color=C["gray"]); c.alignment = al(h="center",v="center")
    ws.row_dimensions[2].height = 22
    r = 4
    r = sec(ws, r, "PARADIGM REFRAME \u2014 INSTRUMENT OF ATTENTION", ncols=7, height=20)
    r = blob(ws, r, (
        "Conventional KPIs measure what is.  Spiral Octave KPIs ask what is alive or asleep in each domain.\n\n"
        "This Songbook is an instrument of attention, not a measurement obligation. "
        "Every cell in the 420 grid is a contemplative position \u2014 a coordinate where asking could happen.\n\n"
        "CEN holds 34 active attentions (the BSC KPIs). "
        "The remaining 386 cells are dormant possibilities of attention \u2014 "
        "not required measurements, not obligations, but a map of where the music could grow."
    ), ncols=7, bg=C["lcyan"], height=18)
    r += 1
    ws.merge_cells(f"A{r}:G{r}")
    c = ws.cell(row=r, column=1, value="THE zeroEnergy PRINCIPLE")
    c.font = fn(sz=11, bold=True, color=C["dcyan"]); ws.row_dimensions[r].height = 18; r += 1
    r = blob(ws, r, (
        '"zeroEnergy: Not absence of effort, but perception below threshold."\n'
        "A zero score does not claim nothing happened. "
        "It claims the instrument\u2019s perception window did not detect whatever did. "
        "This is humility encoded as math."
    ), ncols=7, bg=C["lgray"], height=16, italic=True)
    r += 1
    r = sec(ws, r, "WHAT THIS IS  /  WHAT THIS IS NOT", ncols=7, height=20)
    rows_wt = [
        ("WHAT THIS IS", "WHAT THIS IS NOT"),
        ("A map of 420 coordinates where attention could live in CEN",
         "A list of 420 things CEN must now measure"),
        ("34 BSC KPIs placed at their native dodecahedral coordinates",
         "A replacement for CEN\u2019s Balanced Scorecard"),
        ("A mirror showing which domains are loudly sounding vs. dormant",
         "A verdict on CEN\u2019s performance or capability"),
        ("An instrument of attention \u2014 asks before it measures",
         "An audit or compliance framework"),
        ("A commitment to deliver the full 420 unified workbook in May 2026",
         "A finished product \u2014 Phase 4 is the methodology + 5 examples"),
    ]
    for i, (left, right) in enumerate(rows_wt):
        ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"D{r}:G{r}")
        for col, val in [(1, left), (4, right)]:
            c = ws.cell(row=r, column=col, value=val)
            c.font = fn(bold=(i==0), color=C["white"] if i==0 else (C["dark"] if col==1 else C["red"]))
            c.fill = fl(C["dcyan"] if i==0 else (C["lgreen"] if col==1 else C["lred"]))
            c.alignment = al(h="left", v="top"); c.border = br()
        ws.row_dimensions[r].height = 28; r += 1
    r += 1
    arch = [("Axis","Count","What it is"),
            ("Faces (Domains)","12","Pentagonal faces; organizational capital domains F1-F12"),
            ("Octaves","7","Developmental stages O1 Survival (0.000) to O7 Radiance (0.910 = 1-phi^-5)"),
            ("Elements","5","Pentagramic vertices: Earth, Water, Fire, Air, Ether \u2014 inquiry lenses"),
            ("Total cells","420","Each = (Face x Octave x Element) coordinate \u2014 one contemplative position")]
    for i, (a, b, cv) in enumerate(arch):
        bg = C["dcyan"] if i==0 else C["lcyan"]; fg = C["white"] if i==0 else C["dark"]
        cell(ws, r, 1, a, bold=(i==0), bg=bg, clr=fg, h="center")
        cell(ws, r, 2, b, bold=(i==0), bg=bg, clr=fg, h="center")
        ws.merge_cells(f"C{r}:G{r}")
        c = ws.cell(row=r, column=3, value=cv); c.font = fn(bold=(i==0), color=fg)
        c.fill = fl(bg); c.alignment = al(h="left", v="center"); c.border = br()
        ws.row_dimensions[r].height = 20; r += 1
    ws.merge_cells(f"A{r}:G{r}")
    c = ws.cell(row=r, column=1,
        value="Thesis ref: docs/math/CALCULATION_AUDIT_TRAIL.md  |  SSOT: js/constants/kpi-constants.js:511-524 + phi-harmonics.js:269-347")
    c.font = fn(sz=9, italic=True, color=C["gray"]); c.alignment = al(h="left", v="center")
    cw(ws, [22,22,22,22,14,14,14]); ws.freeze_panes = "A3"
    log.info("Sheet 1 (Cover) built.")

# ── Sheet 2: Methodology ──────────────────────────────────────────────────────
def build_methodology(wb):
    ws = wb.create_sheet("Methodology"); ws.sheet_properties.tabColor = C["purple"]
    ws.merge_cells("A1:F1"); c = ws["A1"]
    c.value = "Mapping Methodology \u2014 Rules A, B, C, Tie-Breakers & Friction"
    c.font = fn(sz=14, bold=True, color=C["dcyan"]); c.alignment = al(h="center",v="center")
    ws.row_dimensions[1].height = 28; r = 3
    r = sec(ws, r, "RULE A \u2014 DOMAIN (Which POC Face F1-F12)")
    r = blob(ws, r, "Read BSC KPI Name + Data Source + Measurement Method. Match to the strongest signal:")
    rule_a = [("Signal in BSC KPI","Primary Face"),
              ("Revenue, costs, P&L, cost coverage","F1  Financial Capital"),
              ("Knowledge, expertise, curriculum, IP","F2  Intellectual Capital"),
              ("People, compensation, roles, hiring, onboarding","F3  Human Capital"),
              ("Policies, governance docs, systems, process","F4  Structural Capital"),
              ("Clients, commercial pipeline, conversion, market","F5  Market Resonance"),
              ("Members, community, partnerships, engagement","F6  Community & Partners"),
              ("NPS, reputation, positioning, visibility","F7  Brand & Reputation"),
              ("Core service delivery, audit throughput, operations","F8  Core Operations"),
              ("Regenerative practice, SDGs, sustainability","F9  Regenerative Flow"),
              ("Values consistency, mission, vision alignment","F10 Foundational Values"),
              ("Donations, subscriptions, fundraising","F11 Funding Pipeline"),
              ("Compliance, GDPR, crisis protocols, independence","F12 Risk & Resilience")]
    for i, (sig, face) in enumerate(rule_a):
        bg = C["dcyan"] if i==0 else (C["lcyan"] if i%2==0 else C["white"])
        fg = C["white"] if i==0 else C["dark"]
        ws.merge_cells(f"A{r}:D{r}"); ws.merge_cells(f"E{r}:F{r}")
        for col, val in [(1,sig),(5,face)]:
            c = ws.cell(row=r, column=col, value=val)
            c.font = fn(bold=(i==0), color=fg); c.fill = fl(bg)
            c.alignment = al(h="left", v="center"); c.border = br()
        ws.row_dimensions[r].height = 17; r += 1
    r += 1
    r = sec(ws, r, "RULE B \u2014 OCTAVE (O1-O7)  |  Use CEN priority tags verbatim")
    r = blob(ws, r, (
        "O1 Survival (Red): 'prerequisites for revenue; CEN cannot scale without these'\n"
        "O2 Structural (Yellow): 'infrastructure for sustainable growth'\n"
        "O3 Aspirational (Green): 'CEN\u2019s vision at full expression'\n"
        "Rule: O1->O1, O2->O2, O3->O3. No re-interpretation. "
        "O4-O7 are the silent higher octaves \u2014 aspirational developmental capacity."), height=18)
    r += 1
    r = sec(ws, r, "RULE C \u2014 ELEMENT (Calibration, Not Invention)")
    prec = [("Face","Domain","Primary Element at O1"),
            ("F1","Financial Capital","Earth"),("F2","Intellectual Capital","Air"),
            ("F3","Human Capital","Fire"),("F4","Structural Capital","Earth"),
            ("F5","Market Resonance","Air"),("F6","Community & Partners","Water"),
            ("F7","Brand & Reputation","Fire"),("F8","Core Operations","Earth"),
            ("F9","Regenerative Flow","Ether"),("F10","Foundational Values","Ether"),
            ("F11","Funding Pipeline","Fire"),("F12","Risk & Resilience","Water")]
    for i, (fa, nm, el) in enumerate(prec):
        bg = C["dcyan"] if i==0 else (C["lcyan"] if i%2==0 else C["white"])
        fg = C["white"] if i==0 else C["dark"]
        for col, val in enumerate([fa,nm,el],1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = fn(bold=(i==0), color=fg); c.fill = fl(bg)
            c.alignment = al(h="center" if col!=2 else "left", v="center"); c.border = br()
        ws.row_dimensions[r].height = 16; r += 1
    r += 1
    esig = [("Element","Inquiry Lens","BSC KPI Type Signal"),
            ("Earth (Foundation)","Material base","Absolute counts, revenue amounts, documents-exist binary, compliance presence"),
            ("Water (Flow)","Circulation","Engagement rate, renewal rate, flow/throughput metrics, cycle time"),
            ("Fire (Catalyst)","Activation","Conversion rate, acquisition rate, ignition mechanisms, capability-operational"),
            ("Air (Connector)","Relationship quality","NPS, touchpoints, quality-of-relationship scores, clarity"),
            ("Ether (Transcendent)","Meaning","Values alignment, vision consistency, purpose integrity, perceived authenticity")]
    for i, (el, lens, sig) in enumerate(esig):
        bg = C["dcyan"] if i==0 else (C["lcyan"] if i%2==0 else C["white"])
        fg = C["white"] if i==0 else C["dark"]
        cell(ws, r, 1, el, bold=(i==0), bg=bg, clr=fg)
        cell(ws, r, 2, lens, bold=(i==0), bg=bg, clr=fg)
        ws.merge_cells(f"C{r}:F{r}")
        c = ws.cell(row=r, column=3, value=sig)
        c.font = fn(bold=(i==0), color=fg); c.fill = fl(bg)
        c.alignment = al(h="left", v="center"); c.border = br()
        ws.row_dimensions[r].height = 17; r += 1
    r += 1
    r = sec(ws, r, "TIE-BREAKER RULES  |  Apply in order")
    for num, rule, detail in [
        ("1","Primary = where measurement LIVES","Whose system generates the number? That is the native face."),
        ("2","Secondary = where IMPACT appears","Noted in Rationale column, not in the coordinate."),
        ("3","One primary assignment per KPI","No dual-primary placements. Dual-impact acknowledged in prose.")]:
        cell(ws,r,1,num,bold=True,bg=C["cyan"],clr=C["white"],h="center")
        ws.merge_cells(f"B{r}:C{r}"); cell(ws,r,2,rule,bold=True,bg=C["lcyan"])
        ws.merge_cells(f"D{r}:F{r}"); cell(ws,r,4,detail,bg=C["lcyan"])
        ws.row_dimensions[r].height = 20; r += 1
    r += 1
    r = sec(ws, r, "RULE FRICTION AS FEATURE")
    blob(ws, r, (
        "If a BSC KPI resists all three rules \u2014 document the friction. Do not hide it. Do not force-fit.\n"
        "Rules that never friction with reality are suspicious. "
        "31 clean + 3 documented-friction mappings is doing the job. "
        "Friction cases surface what the BSC could not say about CEN."), bg=C["lyellow"], italic=True, height=18)
    cw(ws,[15,22,18,26,14,14]); ws.freeze_panes = "A2"
    log.info("Sheet 2 (Methodology) built.")

# ── Sheet 3: 5 Worked Examples ────────────────────────────────────────────────
def build_worked_examples(wb):
    ws = wb.create_sheet("5 Worked Examples"); ws.sheet_properties.tabColor = C["lgreen"]
    ws.merge_cells("A1:G1"); c = ws["A1"]
    c.value = "5 Worked Examples \u2014 BSC KPI to POC Coordinate Mappings"
    c.font = fn(sz=14, bold=True, color=C["dcyan"]); c.alignment = al(h="center",v="center")
    ws.row_dimensions[1].height = 28; r = 3
    cols = ["#","CEN BSC KPI","BSC Perspective / Octave","POC Coordinate\n(Face x Octave x Element)","Precedent","Alignment","Rationale"]
    for i, h in enumerate(cols,1): ws.cell(row=r, column=i, value=h)
    hdr(ws, r, len(cols)); ws.row_dimensions[r].height = 30; r += 1
    examples = [
        ("1","F1  Total annual revenue","Financial / O1 Survival (Red)","F1 Financial Capital\nx O1 Survival\nx Earth","Aligned","Aligned",
         "Revenue as absolute material amount = canonical Earth inquiry at F1. CEN-tagged O1 verbatim. No tie-breaker needed."),
        ("2","C1  Active member count\n(engaged, not ghost)","Customer / O1 Survival (Red)","F6 Community & Partners\nx O1 Survival\nx Earth","Divergent (justified)","Divergent",
         "POC precedent F6->Water (flow). C1 measures a count (Earth), not a flow rate. Member engagement rate (C2) = Water at this cell."),
        ("3","I3  GDPR compliance gaps closed\n(documents-exist binary)","Internal / O1 Survival (Red)","F12 Risk & Resilience\nx O1 Survival\nx Earth","Divergent (justified)","Divergent",
         "POC precedent F12->Water. I3 is a binary (gap closed = Earth). Tie-breaker: F12 not F4 because measurement is risk-specific."),
        ("4","L5  NGCLP curriculum\nready for delivery","Learning & Growth / O3 Aspirational (Green)","F2 Intellectual Capital\nx O3 Relationships\nx Air","Aligned","Aligned",
         "POC precedent F2->Air. Curriculum-ready = knowledge made connectable and teachable. O3 reflects aspirational: not yet delivered."),
        ("5","L8  SDG alignment in PVM\n[CEN\u2019s loudest string: F10 = 8/10]","Learning & Growth / O2 Structural (Yellow)","F10 Foundational Values\nx O2 Structure\nx Ether","Aligned","Aligned",
         "POC precedent F10->Ether. Values-made-structural. F10 V_res_post=8/10. Dominique 10/10, Esther 9/10. 'Your values are your foundation \u2014 and you both know it.'"),
    ]
    for i, (num,kpi,bsc,coord,prec,aln,rat) in enumerate(examples):
        rbg = C["lgray"] if i%2==0 else C["white"]
        abg = C["lgreen"] if aln=="Aligned" else C["lyellow"]
        acl = "00695C" if aln=="Aligned" else C["orange"]
        cell(ws,r,1,num,bold=True,bg=rbg,h="center"); cell(ws,r,2,kpi,bg=rbg)
        cell(ws,r,3,bsc,bg=rbg); cell(ws,r,4,coord,bg=rbg); cell(ws,r,5,prec,bg=rbg)
        cell(ws,r,6,aln,bold=True,bg=abg,clr=acl,h="center"); cell(ws,r,7,rat,bg=rbg)
        ws.row_dimensions[r].height = 55; r += 1
    r += 1
    stats = [("BSC Perspectives","Financial:1, Customer:1, Internal:1, Learning&Growth:2 \u2014 all 4 covered"),
             ("Octaves","O1:3, O2:1, O3:1"),("Elements","Earth:3, Air:1, Ether:1"),
             ("Precedent alignment","Aligned:3, Divergent-justified:2"),
             ("CEN score coverage","Loud strings:1 (F10=8/10)  Quieter:4 (F1=2,F6=4,F12=3,F2=5)")]
    for label, val in stats:
        cell(ws,r,1,label,bold=True,bg=C["lcyan"]); ws.merge_cells(f"B{r}:G{r}")
        c = ws.cell(row=r, column=2, value=val)
        c.font = fn(color=C["dark"]); c.fill = fl(C["lcyan"])
        c.alignment = al(h="left",v="center"); c.border = br()
        ws.row_dimensions[r].height = 17; r += 1
    cw(ws,[4,28,22,22,16,11,42]); ws.freeze_panes = "A4"
    log.info("Sheet 3 (5 Worked Examples) built.")

# ── Sheet 4: Grid Preview (12x7) ──────────────────────────────────────────────
def build_grid_preview(wb, cells_data):
    ws = wb.create_sheet("Grid Preview (12x7)"); ws.sheet_properties.tabColor = C["orange"]
    ws.merge_cells("A1:I1"); c = ws["A1"]
    c.value = "Grid Preview \u2014 12 Faces x 7 Octaves (420 coordinate map)"
    c.font = fn(sz=14, bold=True, color=C["dcyan"]); c.alignment = al(h="center",v="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:I2"); c = ws["A2"]
    c.value = ("Color key:  [Warm orange] F10 Foundational Values \u2014 CEN\u2019s loudest string (8/10)   "
               "[Green] Worked example coordinate   [Blue] Active cells   [Gray] Dormant/silent   "
               "\u2022   Cell notation  [E,W,F,A,E] \u00d75  = all 5 elements (Earth, Water, Fire, Air, Ether) available at this (face \u00d7 octave) coordinate \u2014 each cell holds 5 elemental slots.")
    c.font = fn(sz=9, italic=True, color=C["dark"]); c.alignment = al(h="left",v="center", wrap=True)
    ws.row_dimensions[2].height = 32; r = 3
    # Header
    cell(ws,r,1,"Face",bold=True,bg=C["dcyan"],clr=C["white"],h="center")
    cell(ws,r,2,"Domain Name",bold=True,bg=C["dcyan"],clr=C["white"])
    for oid in range(1,8):
        c = ws.cell(row=r, column=oid+2, value=f"O{oid}\n{OCTAVE[oid]}")
        c.font = fn(bold=True, color=C["white"]); c.fill = fl(C["dcyan"])
        c.alignment = al(h="center",v="center"); c.border = br()
    ws.row_dimensions[r].height = 30; r += 1
    # Build lookup
    lkp = {}
    for cj in cells_data: lkp.setdefault((cj["faceId"],cj["octaveId"]),[]).append(cj)
    for fid in range(1,13):
        name = DOMAIN[fid]
        rbg = C["worange"] if fid==10 else (C["lcyan"] if fid%2==0 else C["white"])
        rfg = C["white"] if fid==10 else C["dark"]
        cell(ws,r,1,f"F{fid}",bold=True,bg=rbg,clr=rfg,h="center")
        cell(ws,r,2,name,bold=(fid==10),bg=rbg,clr=rfg)
        for oid in range(1,8):
            cl_list = lkp.get((fid,oid),[])
            elems = [x["element"][0] for x in cl_list]
            worked = any((fid,oid,x["element"]) in WORKED_COORDS for x in cl_list)
            txt = f"{OCTAVE[oid]}\n[{','.join(elems)}] x{len(cl_list)}"
            if worked:    cbg, cfg = C["lgreen"], "00695C"
            elif fid==10: cbg, cfg = C["lorange"], C["dark"]
            elif cl_list: cbg, cfg = C["lcyan"], C["dark"]
            else:         cbg, cfg = C["lgray"], C["gray"]
            gc = ws.cell(row=r, column=oid+2, value=txt)
            gc.font = fn(sz=8,color=cfg); gc.fill = fl(cbg)
            gc.alignment = al(h="center",v="top"); gc.border = br()
        ws.row_dimensions[r].height = 32; r += 1
    ws.merge_cells(f"A{r}:I{r}")
    c = ws.cell(row=r,column=1,value="F10 Foundational Values (warm): V_res_post=8/10 | Green = worked-example coordinates (Ex1:F1xO1xE, Ex2:F6xO1xE, Ex3:F12xO1xE, Ex4:F2xO3xA, Ex5:F10xO2xEth)")
    c.font = fn(sz=9,italic=True,color=C["orange"]); c.alignment = al(h="left",v="center")
    ws.row_dimensions[r].height = 14
    cw(ws,[6,20,14,14,14,14,14,14,14]); ws.freeze_panes = "A4"
    log.info(f"Sheet 4 (Grid Preview) built. {len(cells_data)} cells.")

# ── Sheet 5: CEN Currently Sounding ──────────────────────────────────────────
def build_currently_sounding(wb):
    ws = wb.create_sheet("CEN Currently Sounding"); ws.sheet_properties.tabColor = C["cyan"]
    ws.merge_cells("A1:F1"); c = ws["A1"]
    c.value = "CEN Currently Sounding \u2014 V_res_post Phase 2 Frozen Scores"
    c.font = fn(sz=14,bold=True,color=C["dcyan"]); c.alignment = al(h="center",v="center")
    ws.row_dimensions[1].height = 28; r = 2
    ws.merge_cells(f"A{r}:F{r}")
    c = ws.cell(row=r,column=1,value="Phase 2 frozen scores \u2014 NOT to be modified. Source: CEN_Phase2_RawScores_Frozen.md. Scale: 1-10 (V_res_post averaged).")
    c.font = fn(sz=9,italic=True,color=C["gray"]); c.alignment = al(h="left",v="center")
    ws.row_dimensions[r].height = 14; r += 2
    for h in ["Face","Domain Name","V_res_post\n(Phase 2)","Volume Category","Color Signal","Notes"]:
        pass
    hdr(ws, r, 6); ws.row_dimensions[r].height = 30
    for i, h in enumerate(["Face","Domain Name","V_res_post\n(Phase 2)","Volume Category","Color Signal","Notes"],1):
        ws.cell(row=r, column=i).value = h
    r += 1
    def vcat(s): return "Loud string" if s>=7 else ("Mid string" if s>=4 else ("Quiet string" if s>=3 else "Dormant"))
    def sbg(s): return C["lgreen"] if s>=7 else (C["lcyan"] if s>=4 else (C["lyellow"] if s>=3 else C["lgray"]))
    f8note = ("Largest co-founder perception gap: D=7, E=1. Per Coherence Portrait: \u2018This is not about who is right. "
              "It is about the fact that CEN\u2019s operational reality looks very different from different positions.\u2019")
    for fid in range(1,13):
        sc = CEN_SCORES[fid]; bg = sbg(sc)
        scl = "00695C" if sc>=7 else (C["dcyan"] if sc>=4 else C["dark"])
        cell(ws,r,1,f"F{fid}",bold=True,bg=bg,h="center"); cell(ws,r,2,DOMAIN[fid],bg=bg)
        c = ws.cell(row=r,column=3,value=sc); c.font = fn(sz=13,bold=True,color=scl)
        c.fill = fl(bg); c.alignment = al(h="center",v="center"); c.border = br()
        cell(ws,r,4,vcat(sc),bg=bg,h="center")
        cell(ws,r,5,("Green" if sc>=7 else "Blue" if sc>=4 else "Yellow" if sc>=3 else "Gray"),bg=bg,h="center")
        cell(ws,r,6,f8note if fid==8 else "",bg=bg)
        ws.row_dimensions[r].height = 45 if fid==8 else 20; r += 1
    r += 1
    for label, faces, interp in [
        ("Loud (7-10)","F9 Regenerative Flow=8, F10 Foundational Values=8","CEN\u2019s strongest strings. The Songbook opens here."),
        ("Mid (4-6)","F2 Intellectual Capital=5; F3=4, F6=4","Active but not at full resonance."),
        ("Quiet (3)","F7 Brand & Reputation=3, F12 Risk & Resilience=3","Beginning to sound."),
        ("Dormant (1-2)","F1=2, F4=2, F5=2, F8=2, F11=2","Below perception threshold. Not failing \u2014 sleeping.")]:
        cell(ws,r,1,label,bold=True,bg=C["lcyan"]); ws.merge_cells(f"B{r}:C{r}")
        c=ws.cell(row=r,column=2,value=faces); c.font=fn(color=C["dark"]); c.fill=fl(C["lcyan"]); c.alignment=al(h="left",v="center"); c.border=br()
        ws.merge_cells(f"D{r}:F{r}")
        c=ws.cell(row=r,column=4,value=interp); c.font=fn(italic=True,color=C["dark"]); c.fill=fl(C["lcyan"]); c.alignment=al(h="left",v="center"); c.border=br()
        ws.row_dimensions[r].height = 20; r += 1
    cw(ws,[8,24,12,16,12,50]); ws.freeze_panes = "A4"
    log.info("Sheet 5 (CEN Currently Sounding) built.")

# ── Sheet 6: Commitment ───────────────────────────────────────────────────────
def build_commitment(wb):
    ws = wb.create_sheet("Commitment"); ws.sheet_properties.tabColor = C["worange"]
    ws.merge_cells("A1:F1"); c = ws["A1"]
    c.value = "Phase 4 Commitment \u2014 May 2026 Delivery Promise"
    c.font = fn(sz=14,bold=True,color=C["dcyan"]); c.alignment = al(h="center",v="center")
    ws.row_dimensions[1].height = 28; r = 3
    r = sec(ws, r, "DELIVERY PROMISE \u2014 MAY 2026")
    r = blob(ws, r, (
        "Unified CEN Reusable Excel Workbook \u2014 extracted from the existing SpiralDashboard architecture:\n\n"
        "  Command Center (Dashboard + control panels)\n"
        "  12 Domain Control Panels (one per face, edge dynamics, vertex vortices)\n"
        "  Full 420-cell Songbook (all 34 BSC KPI mappings, not just 5 examples)\n"
        "  Specification document embedded as reference sheet\n\n"
        "The 420-cell template already exists in SpiralDASHBOARD-Leading_V2.xlsx REFRENCE_MODELS (394/420 cells populated). "
        "Phase 4 = extract + CEN-adapt + deliver as a standalone workbook."), height=18)
    r += 1; r = sec(ws, r, "MILESTONE PLAN")
    hdr(ws, r, 6, bg=C["cyan"]); ws.row_dimensions[r].height = 20
    for i, h in enumerate(["#","Milestone","Target Date","Status","Dependency","Notes"],1):
        ws.cell(row=r, column=i).value = h
    r += 1
    stbg = {"DONE":C["lgreen"],"Pending":C["lyellow"],"Not started":C["lgray"]}
    for num,title,dt,status,dep,note in [
        ("M1","Phase 4 Methodology Songbook delivered to CEN","2026-04-17","DONE","None","This artifact. 5 worked examples + full methodology."),
        ("M2","CEN feedback received on worked examples","2026-04-25","Pending","CEN meeting Apr 17","Which mappings resonate? Which need adjustment?"),
        ("M3","Full 34-KPI mapping completed","2026-05-05","Not started","M2 feedback","29 remaining KPIs after the 5 examples here."),
        ("M4","12 Domain Control Panels built in Excel","2026-05-12","Not started","M3","One panel per face from SpiralDashboard architecture."),
        ("M5","Command Center + full 420 Songbook assembled","2026-05-19","Not started","M3+M4","Unified CEN Reusable Excel Workbook complete."),
        ("M6","Final review + delivery to Dominique & Esther","2026-05-25","Not started","M5","Feeds thesis case study. Defense June 1, 2026.")]:
        rbg = C["lgray"] if int(num[1])%2==0 else C["white"]
        cell(ws,r,1,num,bold=True,bg=rbg,h="center"); cell(ws,r,2,title,bg=rbg)
        cell(ws,r,3,dt,bg=rbg,h="center"); cell(ws,r,4,status,bold=True,bg=stbg.get(status,C["lgray"]),h="center")
        cell(ws,r,5,dep,bg=rbg); cell(ws,r,6,note,bg=rbg)
        ws.row_dimensions[r].height = 22; r += 1
    r += 1; r = sec(ws, r, "3-STEP MEETING USE GUIDE (April 17)")
    for step, title, action in [
        ("Step 1","Open with Cover","Read the Paradigm Reframe together. Establish: instrument of attention, not 420 new KPIs. Let the 'What This Is / What This Is Not' block land."),
        ("Step 2","Show Grid Preview + Currently Sounding","Walk through the 12x7 grid. Point out F10 and F9 as loudest strings. Show the 5 worked examples one at a time. Invite: 'Which mappings feel right? Which feel off?'"),
        ("Step 3","Close with Commitment","Show the milestone plan. Confirm May 2026 delivery of the full 34-KPI mapping and Unified Workbook. Ask: 'What would make this most useful for your ongoing strategy work?'")]:
        cell(ws,r,1,step,bold=True,bg=C["lcyan"],h="center"); ws.merge_cells(f"B{r}:C{r}")
        c=ws.cell(row=r,column=2,value=title); c.font=fn(bold=True,color=C["dcyan"]); c.fill=fl(C["lcyan"]); c.alignment=al(h="left",v="center"); c.border=br()
        ws.merge_cells(f"D{r}:F{r}")
        c=ws.cell(row=r,column=4,value=action); c.font=fn(color=C["dark"]); c.fill=fl(C["lcyan"]); c.alignment=al(h="left",v="top"); c.border=br()
        ws.row_dimensions[r].height = 42; r += 1
    cw(ws,[10,30,14,12,20,38]); ws.freeze_panes = "A3"
    log.info("Sheet 6 (Commitment) built.")

# ── Sheet 7: Thesis Contribution ─────────────────────────────────────────────
def build_thesis_contribution(wb, spiral_path=None):
    ws = wb.create_sheet("Thesis Contribution"); ws.sheet_properties.tabColor = C["purple"]
    ws.merge_cells("A1:F1"); c = ws["A1"]
    c.value = "Thesis Contribution \u2014 SQ5 & Spiral Quality Scores"
    c.font = fn(sz=14,bold=True,color=C["dcyan"]); c.alignment = al(h="center",v="center")
    ws.row_dimensions[1].height = 28; r = 3
    r = sec(ws, r, "RESEARCH QUESTION SQ5 CLAIM")
    r = blob(ws, r, (
        "\u201cConventional KPIs measure what is. Spiral Octave KPIs ask what is alive or asleep in each domain.\u201d\n\n"
        "This claim operationalizes the Spiral Octave Songbook as an epistemological upgrade to the BSC \u2014 not a replacement. "
        "Every BSC KPI is absorbed into the 420 framework at its native coordinate. "
        "The 386 remaining cells = aspirational developmental capacity, not obligations.\n\n"
        "The upgrade is from measurement to contemplative inquiry: numbers attached to questions, not abstract dashboards."), height=18)
    r += 1
    ws.merge_cells(f"A{r}:F{r}")
    ws.cell(row=r,column=1,value="THESIS CROSS-REFERENCES").font = fn(sz=11,bold=True,color=C["dcyan"])
    ws.row_dimensions[r].height = 18; r += 1
    hdr(ws, r, 3)
    for i, h in enumerate(["Reference","Path","Purpose"],1): ws.cell(row=r,column=i,value=h)
    ws.merge_cells(f"C{r}:F{r}")
    ws.cell(row=r,column=3).font = fn(bold=True,color=C["white"]); ws.cell(row=r,column=3).fill = fl(C["dcyan"])
    ws.row_dimensions[r].height = 18; r += 1
    for i, (ref, path, purp) in enumerate([
        ("Calculation Audit Trail","docs/math/CALCULATION_AUDIT_TRAIL.md","Every formula traced input to output. Academic armor for the defense."),
        ("Mathematical Audit","Final Thesis/.../SpiralDASHBOARD_V2_Full_Mathematical_Audit.md","2,022 formulas documented. Honest 60/420 operational scope disclosure."),
        ("420 Canonical Template","Final Thesis/.../SpiralDASHBOARD-Leading_V2.xlsx (REFRENCE_MODELS)","All 420 cells defined. 394/420 with values. This Songbook extends it to CEN."),
        ("CEN Phase 2 Frozen Scores","Final Thesis/.../CEN_Phase2_RawScores_Frozen.md","V_res_post canonical scores. Sheet 5 hardcodes these verbatim."),
        ("Emergent Mathematical Truths","docs/thesis/EMERGENT_MATHEMATICAL_TRUTHS.md","11 discoveries. #7,#10,#11: phi resonance with dodecahedral topology.")]):
        bg = C["lcyan"] if i%2==0 else C["white"]
        cell(ws,r,1,ref,bold=True,bg=bg); ws.merge_cells(f"B{r}:C{r}")
        c=ws.cell(row=r,column=2,value=path); c.font=fn(sz=9,italic=True,color=C["dark"]); c.fill=fl(bg); c.alignment=al(h="left",v="center"); c.border=br()
        ws.merge_cells(f"D{r}:F{r}")
        c=ws.cell(row=r,column=4,value=purp); c.font=fn(color=C["dark"]); c.fill=fl(bg); c.alignment=al(h="left",v="center"); c.border=br()
        ws.row_dimensions[r].height = 20; r += 1
    r += 1; r = sec(ws, r, "SPIRAL REFINEMENT LOOP QUALITY REPORT")
    if spiral_path:
        scores = _parse_spiral_report(spiral_path)
        if scores:
            ws.merge_cells(f"A{r}:F{r}")
            ws.cell(row=r,column=1,value=f"Source: {spiral_path}").font = fn(sz=9,italic=True,color=C["gray"])
            ws.row_dimensions[r].height = 14; r += 1
            hdr(ws, r, 5, bg=C["purple"]); ws.row_dimensions[r].height = 20
            for i, h in enumerate(["Face #","Face Name","Score (1-10)","Floor (9/10)","Status"],1):
                ws.cell(row=r,column=i).value = h
            r += 1
            for fn_num, (fname, sc) in sorted(scores.items()):
                passed = sc >= 9; bg = C["lgreen"] if passed else C["lred"]
                status = "PASS" if passed else "BELOW FLOOR"
                cell(ws,r,1,fn_num,bold=True,bg=bg,h="center")
                cell(ws,r,2,fname,bg=bg)
                c = ws.cell(row=r,column=3,value=sc)
                c.font = fn(sz=12,bold=True,color="00695C" if passed else C["red"])
                c.fill = fl(bg); c.alignment = al(h="center",v="center"); c.border = br()
                cell(ws,r,4,"9/10",bg=bg,h="center"); cell(ws,r,5,status,bold=True,bg=bg,clr="00695C" if passed else C["red"],h="center")
                ws.row_dimensions[r].height = 20; r += 1
        else:
            log.warning("Spiral report found but no scores extracted \u2014 using placeholder.")
            _spiral_placeholder(ws, r, spiral_path)
    else:
        _spiral_placeholder(ws, r)
    cw(ws,[12,28,14,14,16]); ws.freeze_panes = "A3"
    log.info("Sheet 7 (Thesis Contribution) built.")

def _spiral_placeholder(ws, row, path=None):
    ws.merge_cells(f"A{row}:F{row+2}")
    msg = ("Spiral report pending \u2014 will be embedded when spiral refinement loop completes.\n\n"
           "Run: python _build_songbook_phase4_commitment.py --spiral-report PATH\n"
           "where PATH is the markdown file containing the 12-face spiral score table.")
    if path: msg = f"[WARN] Could not extract scores from: {path}\n\n" + msg
    c = ws.cell(row=row, column=1, value=msg)
    c.font = fn(sz=10,italic=True,color=C["gray"]); c.fill = fl(C["lgray"])
    c.alignment = al(h="left",v="top"); c.border = br()
    for rr in range(row, row+3): ws.row_dimensions[rr].height = 18

def _parse_spiral_report(path):
    FACES = ["Survival","Foundation","Clarity","Harmony","Signal","Consciousness",
             "Coherence","Resilience","Growth","Integrity","Discovery","Radiance"]
    try:
        content = open(path, encoding="utf-8").read()
    except Exception as e:
        log.error(f"Cannot read spiral report: {e}"); return {}
    scores = {}
    for m in re.finditer(r"\|\s*(\d{1,2})\s*\|\s*([A-Za-z &]+?)\s*\|\s*(\d+(?:\.\d+)?)\s*\|", content):
        fn_num, fname, sc = int(m.group(1)), m.group(2).strip(), float(m.group(3))
        if 1<=fn_num<=12 and 1<=sc<=10: scores[fn_num] = (fname, sc)
    if not scores:
        for m in re.finditer(r"\|\s*("+'|'.join(FACES)+r")\s*\|\s*(\d+(?:\.\d+)?)\s*\|", content, re.IGNORECASE):
            name = m.group(1).strip().title(); sc = float(m.group(2))
            if name in FACES and 1<=sc<=10: scores[FACES.index(name)+1] = (name, sc)
    log.info(f"Parsed {len(scores)} face scores from spiral report." if scores else "No face scores found.")
    return scores

# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Build POC_420_Songbook_Phase4_Commitment.xlsx")
    parser.add_argument("--spiral-report", metavar="PATH",
                        help="Markdown spiral report for Sheet 7 face scores.")
    args = parser.parse_args()
    log.info("Starting Songbook Phase 4 Commitment build.")
    log.info(f"JSON input: {JSON_PATH}"); log.info(f"Output: {OUT_PATH}")
    if not JSON_PATH.exists():
        log.error(f"JSON not found: {JSON_PATH}"); sys.exit(1)
    with open(JSON_PATH, encoding="utf-8") as f: raw = json.load(f)
    cells_data = raw.get("cells", [])
    log.info(f"Loaded {len(cells_data)} cells from JSON.")
    wb = openpyxl.Workbook()
    build_cover(wb); build_methodology(wb); build_worked_examples(wb)
    build_grid_preview(wb, cells_data); build_currently_sounding(wb)
    build_commitment(wb); build_thesis_contribution(wb, args.spiral_report)
    log.info(f"Writing to temp: {TMP_PATH}")
    try: wb.save(str(TMP_PATH))
    except Exception as e: log.error(f"Save failed: {e}"); sys.exit(1)
    # Validate
    try:
        twb = openpyxl.load_workbook(str(TMP_PATH))
        expected = ["Cover","Methodology","5 Worked Examples","Grid Preview (12x7)",
                    "CEN Currently Sounding","Commitment","Thesis Contribution"]
        missing = [s for s in expected if s not in twb.sheetnames]
        if missing: log.error(f"Missing sheets: {missing}"); sys.exit(1)
        cover_txt = " ".join(str(c.value) for row in twb["Cover"].iter_rows() for c in row if c.value)
        for req in ["instrument of attention","zeroEnergy"]:
            if req not in cover_txt: log.error(f"Smoke test: '{req}' not in Cover"); sys.exit(1)
        if not args.spiral_report:
            t7 = " ".join(str(c.value) for row in twb["Thesis Contribution"].iter_rows() for c in row if c.value)
            if "spiral report pending" not in t7.lower():
                log.error("Smoke test: 'spiral report pending' not in Sheet 7"); sys.exit(1)
        log.info(f"Validation passed. Sheets: {twb.sheetnames}")
    except SystemExit: raise
    except Exception as e: log.error(f"Validation error: {e}"); sys.exit(1)
    if OUT_PATH.exists(): OUT_PATH.unlink()
    TMP_PATH.rename(OUT_PATH)
    size_kb = OUT_PATH.stat().st_size // 1024
    log.info(f"Done. {OUT_PATH} ({size_kb} KB)")
    print(f"[SUCCESS] {OUT_PATH}  ({size_kb} KB)  sheets={len(expected)}")

if __name__ == "__main__":
    main()
